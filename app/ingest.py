"""知识库导入与索引构建。

流程：读 JSON → 写 standards/clauses 主表 → 重建全文索引（jieba 分词）
      → 重建语义向量 → 递增索引版本号（让检索缓存失效）。

任何一步失败都不应该让整个知识库不可用：
语义模型下载失败时，只跳过向量这一步，关键词检索照常用。
"""
from __future__ import annotations

import json
import logging
import re
from pathlib import Path
from typing import Callable, Iterable

from . import config, db
from .normalize import build_search_text

ProgressCb = Callable[[str, int, int], None]

# 建向量时每条的文本：标题 + 摘要 + 关键词 + 正文（截断）
EMBED_TEXT_LIMIT = 400


def _noop(stage: str, done: int, total: int) -> None:
    return None


# --------------------------------------------------------------------------
# 标准与条款写入
# --------------------------------------------------------------------------
def upsert_standard(meta: dict) -> str:
    std_id = str(meta.get("id") or "").strip()
    if not std_id:
        raise ValueError("standard.id 不能为空")
    db.execute(
        """
        INSERT INTO standards(id, name, full_name, version, publisher, license, source_url, note, imported_at)
        VALUES(?,?,?,?,?,?,?,?,?)
        ON CONFLICT(id) DO UPDATE SET
            name=excluded.name, full_name=excluded.full_name, version=excluded.version,
            publisher=excluded.publisher, license=excluded.license,
            source_url=excluded.source_url, note=excluded.note,
            imported_at=excluded.imported_at
        """,
        (
            std_id,
            meta.get("name") or std_id,
            meta.get("full_name") or "",
            meta.get("version") or "",
            meta.get("publisher") or "",
            meta.get("license") or "",
            meta.get("source_url") or "",
            meta.get("note") or "",
            db.now(),
        ),
    )
    db.set_setting("standards_version", db.now())
    return std_id


def _normalize_clause(raw: dict, standard_id: str) -> dict | None:
    clause_id = str(raw.get("clause_id") or "").strip()
    title = str(raw.get("title") or "").strip()
    text = str(raw.get("text") or "").strip()
    if not clause_id or not title:
        return None
    text_type = str(raw.get("text_type") or "summary").strip().lower()
    if text_type not in {"original", "summary"}:
        text_type = "summary"

    chapter_path = raw.get("chapter_path") or []
    if isinstance(chapter_path, str):
        chapter_path = [chapter_path]

    def _as_list(value) -> list[str]:
        if value is None:
            return []
        if isinstance(value, str):
            return [value]
        return [str(v) for v in value if str(v).strip()]

    return {
        "standard_id": standard_id,
        "clause_id": clause_id,
        "title": title,
        "chapter_path": db.dump_json(_as_list(chapter_path)),
        "text": text or title,
        "text_type": text_type,
        "summary": str(raw.get("summary") or "").strip(),
        "keywords": db.dump_json(_as_list(raw.get("keywords"))),
        "risk_tags": db.dump_json(_as_list(raw.get("risk_tags"))),
        "level_scope": db.dump_json(_as_list(raw.get("level_scope"))),
        "sort_key": str(raw.get("sort_key") or _default_sort_key(clause_id)).strip(),
    }


def _default_sort_key(clause_id: str) -> str:
    """把 8.1.4.10 这类编号补零，保证按章节正确排序。"""
    parts = re.split(r"([0-9]+)", clause_id)
    out = []
    for p in parts:
        if p.isdigit():
            out.append(p.zfill(4))
        else:
            out.append(p)
    return "".join(out)


def upsert_clauses(
    standard_id: str,
    clauses: Iterable[dict],
    *,
    source_kind: str = "builtin",
    batch_id: int | None = None,
) -> tuple[int, int, int]:
    """写入条款。返回 (新增, 更新, 跳过)。

    source_kind / batch_id 用来追溯来源：内置种子还是用户导入、来自哪一批。
    """
    added = updated = skipped = 0
    rows = []
    for raw in clauses:
        item = _normalize_clause(raw, standard_id)
        if item is None:
            skipped += 1
            continue
        rows.append(item)

    if not rows:
        return 0, 0, skipped

    existing = {
        r["clause_id"]
        for r in db.query("SELECT clause_id FROM clauses WHERE standard_id = ?", (standard_id,))
    }

    with db.get_conn() as conn:
        for item in rows:
            is_new = item["clause_id"] not in existing
            conn.execute(
                """
                INSERT INTO clauses(standard_id, clause_id, title, chapter_path, text, text_type,
                                    summary, keywords, risk_tags, level_scope, sort_key, enabled,
                                    created_at, source_kind, batch_id, imported_at)
                VALUES(:standard_id,:clause_id,:title,:chapter_path,:text,:text_type,
                       :summary,:keywords,:risk_tags,:level_scope,:sort_key,1,
                       :created_at,:source_kind,:batch_id,:imported_at)
                ON CONFLICT(standard_id, clause_id) DO UPDATE SET
                    title=excluded.title, chapter_path=excluded.chapter_path, text=excluded.text,
                    text_type=excluded.text_type, summary=excluded.summary,
                    keywords=excluded.keywords, risk_tags=excluded.risk_tags,
                    level_scope=excluded.level_scope, sort_key=excluded.sort_key,
                    source_kind=excluded.source_kind, batch_id=excluded.batch_id,
                    imported_at=excluded.imported_at
                """,
                {
                    **item,
                    "created_at": db.now(),
                    "source_kind": source_kind,
                    "batch_id": batch_id,
                    "imported_at": db.now(),
                },
            )
            if is_new:
                added += 1
            else:
                updated += 1
    return added, updated, skipped


# --------------------------------------------------------------------------
# 导入批次（来源追溯 / 撤销用）
# --------------------------------------------------------------------------
def _create_batch(filename: str, standard_id: str, standard_name: str,
                  kind: str, source_kind: str) -> int:
    return db.execute(
        "INSERT INTO import_batches(filename, standard_id, standard_name, kind, source_kind, created_at)"
        " VALUES(?,?,?,?,?,?)",
        (filename, standard_id, standard_name, kind, source_kind, db.now()),
    )


def _finish_batch(batch_id: int, added: int, updated: int, skipped: int, note: str = "") -> None:
    db.execute(
        "UPDATE import_batches SET added = ?, updated = ?, skipped = ?, note = ? WHERE id = ?",
        (added, updated, skipped, note, batch_id),
    )


def restore_batch_record(row: dict, *, reverted: int, note: str) -> int:
    """把一条批次记录重新写回数据库。

    为什么需要这一步：
    备份是在**创建批次记录之前**做的，所以恢复备份会把这条批次记录一起抹掉。
    如果不补回来，用户看不到"这批被撤销过"，重复撤销的拦截也会失效（变成 404）。
    """
    db.execute(
        "INSERT OR REPLACE INTO import_batches"
        "(id, filename, standard_id, standard_name, kind, source_kind,"
        " added, updated, skipped, note, backup_path, reverted, created_at)"
        " VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (
            row["id"], row["filename"], row["standard_id"], row["standard_name"],
            row["kind"], row["source_kind"], row["added"], row["updated"],
            row["skipped"], note, row["backup_path"], reverted, row["created_at"],
        ),
    )
    return int(row["id"])


# --------------------------------------------------------------------------
# 索引重建
# --------------------------------------------------------------------------
def rebuild_fts(progress: ProgressCb = _noop) -> int:
    """重建全文索引 + 词频统计（IDF 用）。

    tokens 列存 jieba 分词结果，rowid 与 clauses.id 对齐。
    token_df 表记录每个词出现在多少条条款里，检索时据此算 IDF 权重。
    """
    from collections import Counter

    rows = db.query(
        "SELECT id, clause_id, title, chapter_path, text, summary, keywords FROM clauses"
    )
    total = len(rows)
    df_counter: Counter = Counter()

    with db.get_conn() as conn:
        conn.execute("DELETE FROM clauses_fts")
        for i, r in enumerate(rows, 1):
            keywords = " ".join(db.load_json(r["keywords"]))
            chapter = " ".join(db.load_json(r["chapter_path"]))
            tokens = build_search_text(
                r["clause_id"], r["title"], chapter, keywords, r["summary"], r["text"]
            )
            conn.execute(
                "INSERT INTO clauses_fts(rowid, tokens) VALUES (?, ?)", (r["id"], tokens)
            )
            # 同一条条款里的重复词只算一次（文档频率，不是词频）
            for token in set(tokens.split()):
                df_counter[token] += 1
            if i % 50 == 0 or i == total:
                progress("全文索引", i, total)

        conn.execute("DELETE FROM token_df")
        conn.executemany(
            "INSERT INTO token_df(token, df) VALUES (?, ?)",
            list(df_counter.items()),
        )
    return total


def embed_text_for(clause: dict) -> str:
    parts = [
        clause.get("title") or "",
        clause.get("summary") or "",
        " ".join(clause.get("keywords") or []),
        " ".join(clause.get("chapter_path") or []),
        (clause.get("text") or "")[:EMBED_TEXT_LIMIT],
    ]
    return " ".join(p for p in parts if p).strip()


def rebuild_vectors(progress: ProgressCb = _noop, batch_size: int = 32) -> tuple[int, str]:
    """重建语义向量。返回 (写入条数, 说明)。模型不可用时返回 (0, 原因)。"""
    from .retrieval import get_engine

    engine = get_engine()
    if not config.ENABLE_VECTOR_SEARCH:
        return 0, "配置里已关闭语义检索（ENABLE_VECTOR_SEARCH=false）"
    if engine._get_embedder() is None:
        return 0, "语义模型不可用（未下载成功或依赖缺失）"

    rows = db.query("SELECT * FROM clauses WHERE enabled = 1")
    total = len(rows)
    if total == 0:
        return 0, "没有条款需要向量化"

    done = 0
    with db.get_conn() as conn:
        conn.execute("DELETE FROM clause_vectors")
        buffer_ids: list[int] = []
        buffer_texts: list[str] = []

        def flush() -> None:
            nonlocal done, buffer_ids, buffer_texts
            if not buffer_ids:
                return
            vectors = engine.embed(buffer_texts)
            for rowid, vec in zip(buffer_ids, vectors):
                vec = [float(x) for x in vec]
                conn.execute(
                    "INSERT OR REPLACE INTO clause_vectors(clause_rowid, dim, vec) VALUES (?,?,?)",
                    (rowid, len(vec), db_blob(vec)),
                )
            done += len(buffer_ids)
            buffer_ids, buffer_texts = [], []
            progress("语义向量", done, total)

        for r in rows:
            clause = db.clause_to_dict(r)
            buffer_ids.append(int(r["id"]))
            buffer_texts.append(embed_text_for(clause))
            if len(buffer_ids) >= batch_size:
                flush()
        flush()

    return done, ""


def db_blob(vec: list[float]) -> bytes:
    import struct

    return struct.pack(f"<{len(vec)}f", *vec)


def rebuild_index(progress: ProgressCb = _noop, with_vectors: bool = True) -> dict:
    """完整重建索引。"""
    n_fts = rebuild_fts(progress)
    n_vec, note = (0, "")
    if with_vectors:
        n_vec, note = rebuild_vectors(progress)
    db.set_setting("index_version", db.now())
    db.set_setting("index_clauses", str(n_fts))
    db.set_setting("index_vectors", str(n_vec))

    from .retrieval import reset_engine_cache

    reset_engine_cache()
    return {"fts": n_fts, "vectors": n_vec, "vector_note": note}


# --------------------------------------------------------------------------
# 种子知识库
# --------------------------------------------------------------------------
def import_standard_file(
    path: Path,
    progress: ProgressCb = _noop,
    *,
    kind: str = "seed",
    source_kind: str = "builtin",
) -> dict:
    """导入标准 JSON 文件。

    kind:        seed（内置种子）/ json（用户上传的 JSON）
    source_kind: builtin / user —— 决定界面上显示的来源标识
    """
    with open(path, "r", encoding="utf-8") as f:
        data = json.load(f)

    meta = data.get("standard") or {}
    clauses = data.get("clauses") or []
    if not isinstance(clauses, list):
        raise ValueError(f"{path.name} 里 clauses 不是数组")

    std_id = upsert_standard(meta)
    std_name = meta.get("name") or std_id
    batch_id = _create_batch(path.name, std_id, std_name, kind, source_kind)

    added, updated, skipped = upsert_clauses(
        std_id, clauses, source_kind=source_kind, batch_id=batch_id
    )
    note = f"跳过 {skipped} 条" if skipped else ""
    _finish_batch(batch_id, added, updated, skipped, note)
    db.execute(
        "INSERT INTO import_log(filename, standard_id, clauses_added, clauses_updated, note, created_at)"
        " VALUES(?,?,?,?,?,?)",
        (path.name, std_id, added, updated, note, db.now()),
    )
    return {
        "filename": path.name,
        "standard_id": std_id,
        "standard_name": std_name,
        "batch_id": batch_id,
        "source_kind": source_kind,
        "added": added,
        "updated": updated,
        "skipped": skipped,
    }


def _is_standard_file(path: Path) -> bool:
    """判断是不是"标准条款"文件。

    data/seed 下还放着 risk_terms.json（风险词表）这类配置文件，
    它们没有 standard.id，不能被当成标准导入。
    """
    if path.name.startswith("_") or path.suffix.lower() != ".json":
        return False
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        meta = data.get("standard") if isinstance(data, dict) else None
        return isinstance(meta, dict) and bool(str(meta.get("id") or "").strip())
    except Exception:
        return False


def ensure_seed_loaded(progress: ProgressCb = _noop) -> dict:
    """首次启动自动导入种子知识库（已有条款且版本一致就跳过）。"""
    from .retrieval import get_engine

    seed_files = [p for p in sorted(config.SEED_DIR.glob("*.json")) if _is_standard_file(p)]
    if not seed_files:
        return {"imported": [], "skipped": True, "reason": "data/seed 下没有可导入的标准文件"}

    current = int(db.query_one("SELECT COUNT(*) AS c FROM clauses")["c"])
    stamp = db.get_setting("seed_signature", "")
    signature = "|".join(f"{p.name}:{int(p.stat().st_mtime)}:{p.stat().st_size}" for p in seed_files)

    if current > 0 and stamp == signature:
        return {"imported": [], "skipped": True, "reason": f"种子知识库已是最新（{current} 条条款）"}

    results = []
    for path in seed_files:
        try:
            results.append(import_standard_file(path, progress))
        except Exception as exc:
            results.append({"filename": path.name, "error": str(exc)})

    rebuilt = rebuild_index(progress)
    db.set_setting("seed_signature", signature)
    # 预加载模型，避免用户第一次提问时干等
    try:
        get_engine().warmup()
    except Exception as exc:
        logging.getLogger(__name__).warning("语义模型预热失败：%s", exc)
    return {"imported": results, "skipped": False, "rebuilt": rebuilt}


# --------------------------------------------------------------------------
# 导入用户自己的文档（PDF / Word / Markdown / TXT）
# --------------------------------------------------------------------------
# 识别"条款标题行"的几种常见写法
HEADING_PATTERNS = [
    # 8.1.4.1 身份鉴别   /  8.1.4.1、身份鉴别
    re.compile(r"^\s*(\d+(?:\.\d+){1,4})\s*[、.．:：]?\s*(.{2,50})$"),
    # 第8章 / 第三条 / 第一节
    re.compile(r"^\s*(第\s*[0-9一二三四五六七八九十百]+\s*[章节条款])\s*[、.．:：]?\s*(.{0,50})$"),
    # 一、身份鉴别  /  （一）身份鉴别
    re.compile(r"^\s*[（(]?\s*([一二三四五六七八九十]+)\s*[）)]?\s*[、.．]\s*(.{2,50})$"),
    # 1. 身份鉴别 / 1、身份鉴别
    re.compile(r"^\s*(\d{1,3})\s*[、.．]\s*(.{2,50})$"),
]

MIN_CLAUSE_CHARS = 15
MAX_CLAUSE_CHARS = 1200


def _read_document(path: Path) -> list[str]:
    """把文档读成按行拆分的文本。"""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("缺少 pypdf，无法解析 PDF，请先安装依赖") from exc
        reader = PdfReader(str(path))
        lines: list[str] = []
        for page in reader.pages:
            text = page.extract_text() or ""
            lines.extend(text.splitlines())
        return lines

    if suffix in {".docx", ".doc"}:
        try:
            import docx  # python-docx
        except ImportError as exc:
            raise RuntimeError("缺少 python-docx，无法解析 Word，请先安装依赖") from exc
        if suffix == ".doc":
            raise RuntimeError("暂不支持旧版 .doc，请先用 Word 另存为 .docx")
        document = docx.Document(str(path))
        lines = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    lines.append(" ".join(cells))
        return lines

    # .md / .txt / 其他纯文本
    for encoding in ("utf-8", "utf-8-sig", "gbk", "gb18030"):
        try:
            return path.read_text(encoding=encoding).splitlines()
        except (UnicodeDecodeError, LookupError):
            continue
    return path.read_text(encoding="utf-8", errors="ignore").splitlines()


def split_into_clauses(lines: list[str], prefix: str) -> list[dict]:
    """按标题行把文档切成条款。切不出标题时退化为定长分段。"""
    clauses: list[dict] = []
    current_id: str | None = None
    current_title = ""
    buffer: list[str] = []
    counter = 0

    def flush() -> None:
        nonlocal current_id, current_title, buffer
        body = "\n".join(x for x in buffer if x.strip()).strip()
        if current_id and len(body) >= MIN_CLAUSE_CHARS:
            clauses.append({
                "clause_id": current_id,
                "title": current_title or current_id,
                "chapter_path": [prefix, "导入内容"],
                "text": body[:MAX_CLAUSE_CHARS],
                "text_type": "original",
                "summary": body[:60].replace("\n", " "),
                "keywords": [],
                "risk_tags": [],
                "level_scope": [],
                "sort_key": current_id,
            })
        current_id, current_title, buffer = None, "", []

    for raw_line in lines:
        line = raw_line.strip().lstrip("#").strip()
        if not line:
            if buffer:
                buffer.append("")
            continue

        matched = None
        for pattern in HEADING_PATTERNS:
            m = pattern.match(line)
            if m:
                matched = m
                break

        if matched:
            flush()
            counter += 1
            number = matched.group(1).strip()
            title = (matched.group(2) or "").strip() or number
            # 纯数字序号（如 "1. 身份鉴别"）加前缀避免跨文档冲突
            current_id = number if "." in number or "第" in number else f"{prefix}-{number}"
            current_title = title[:60]
        else:
            if current_id is None:
                counter += 1
                current_id = f"{prefix}-{counter:03d}"
                current_title = line[:60]
                buffer = []
            buffer.append(line)

    flush()

    # 一条都没切出来 → 定长分段兜底
    if not clauses:
        text = "\n".join(x for x in lines if x.strip()).strip()
        for i in range(0, len(text), MAX_CLAUSE_CHARS):
            chunk = text[i : i + MAX_CLAUSE_CHARS].strip()
            if len(chunk) < MIN_CLAUSE_CHARS:
                continue
            n = i // MAX_CLAUSE_CHARS + 1
            clauses.append({
                "clause_id": f"{prefix}-{n:03d}",
                "title": chunk[:40].replace("\n", " "),
                "chapter_path": [prefix, "导入内容"],
                "text": chunk,
                "text_type": "original",
                "summary": chunk[:60].replace("\n", " "),
                "keywords": [],
                "risk_tags": [],
                "level_scope": [],
                "sort_key": f"{prefix}-{n:03d}",
            })
    return clauses


def import_document(path: Path, progress: ProgressCb = _noop) -> dict:
    """导入用户自己的 PDF / Word / Markdown / TXT 文档。"""
    lines = _read_document(path)
    if not lines:
        raise RuntimeError("文档里没有读到任何文字（可能是扫描版 PDF，需要先做 OCR）")

    # 用文件名做标准 id
    stem = re.sub(r"[^0-9A-Za-z\u4e00-\u9fff]+", "_", path.stem).strip("_") or "imported"
    prefix = f"IMP-{stem[:16]}"
    std_id = f"custom_{stem[:32]}"

    clauses = split_into_clauses(lines, prefix)
    if not clauses:
        raise RuntimeError("没能从文档里切分出有效条款，请检查文件内容")

    meta = {
        "id": std_id,
        "name": path.stem[:40],
        "full_name": path.name,
        "version": "用户导入",
        "publisher": "用户自备资料",
        "license": "用户自备",
        "source_url": "",
        "note": "由用户自行导入的资料，条款编号与切分由程序自动识别，可能不精确，请自行核对。",
    }
    upsert_standard(meta)
    batch_id = _create_batch(path.name, std_id, path.stem[:40], "document", "user")
    added, updated, skipped = upsert_clauses(
        std_id, clauses, source_kind="user", batch_id=batch_id
    )
    note = f"自动切分，跳过 {skipped} 段"
    _finish_batch(batch_id, added, updated, skipped, note)
    db.execute(
        "INSERT INTO import_log(filename, standard_id, clauses_added, clauses_updated, note, created_at)"
        " VALUES(?,?,?,?,?,?)",
        (path.name, std_id, added, updated, note, db.now()),
    )
    return {
        "filename": path.name,
        "standard_id": std_id,
        "standard_name": path.stem[:40],
        "batch_id": batch_id,
        "source_kind": "user",
        "added": added,
        "updated": updated,
        "skipped": skipped,
        "note": f"已从文档中自动切分出 {len(clauses)} 条，编号为程序生成，请核对。",
    }
